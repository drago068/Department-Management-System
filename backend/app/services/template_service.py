import io
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def generate_student_enrollment_template() -> io.BytesIO:
    """Generates an official styled Excel template (.xlsx) for bulk student enrollment."""
    wb = openpyxl.Workbook()

    # ─────────────────────────────────────────────────────────────
    # SHEET 1: Student Roster Template
    # ─────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Student_Roster"
    ws.views.sheetView[0].showGridLines = True

    headers = [
        "register_number",
        "full_name",
        "email",
        "phone",
        "parent_name",
        "parent_phone",
        "gender",
        "date_of_birth",
        "blood_group",
        "address",
        "department_code",
        "batch_name",
        "section_name",
    ]

    header_labels = [
        "Register Number *",
        "Full Name *",
        "Email (Optional)",
        "Mobile Number",
        "Parent Name",
        "Parent Phone",
        "Gender",
        "Date of Birth (YYYY-MM-DD)",
        "Blood Group",
        "Address",
        "Department Code",
        "Batch Name",
        "Section",
    ]

    # Color tokens
    header_fill = PatternFill(start_color="003FB1", end_color="003FB1", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    # Write Header Row
    ws.row_dimensions[1].height = 28
    for col_idx, (col_id, label) in enumerate(zip(headers, header_labels), start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_id)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    # Sample Data Rows
    sample_data = [
        [
            "714022AD001",
            "Arjun Patel",
            "arjun.p@suguna.edu",
            "+91 9876543210",
            "Rajesh Patel",
            "+91 9876500000",
            "Male",
            "2004-05-14",
            "O+",
            "142, North Street, Coimbatore",
            "AIDS",
            "2024-2028",
            "A",
        ],
        [
            "714022AD002",
            "Bhavana S",
            "bhavana.s@suguna.edu",
            "+91 9876543211",
            "Suresh S",
            "+91 9876500001",
            "Female",
            "2004-08-22",
            "A+",
            "54, Gandhi Nagar, Coimbatore",
            "AIDS",
            "2024-2028",
            "A",
        ],
        [
            "714022AD003",
            "Rahul V",
            "rahul.v@suguna.edu",
            "+91 9876543212",
            "Vijay Kumar",
            "+91 9876500002",
            "Male",
            "2004-11-10",
            "B+",
            "12, Crosscut Road, Coimbatore",
            "AIDS",
            "2024-2028",
            "A",
        ],
        [
            "24AD012",
            "Abishek R K",
            "abishek.24ad012@suguna.edu.in",
            "+91 9842177310",
            "K. Rangarajan",
            "+91 9443210892",
            "Male",
            "2005-04-18",
            "O+",
            "42, Bharathi Park 7th Cross, Coimbatore",
            "AIDS",
            "2024-2028",
            "A",
        ],
    ]

    sample_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    data_font = Font(name="Calibri", size=10, color="1F2937")

    for row_idx, row_values in enumerate(sample_data, start=2):
        ws.row_dimensions[row_idx].height = 22
        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.alignment = left_align
            cell.border = thin_border
            if row_idx % 2 == 1:
                cell.fill = sample_fill

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    # ─────────────────────────────────────────────────────────────
    # SHEET 2: Instructions & Predefined Password Rules
    # ─────────────────────────────────────────────────────────────
    ws_guide = wb.create_sheet(title="Instructions_&_Rules")
    ws_guide.views.sheetView[0].showGridLines = True

    title_font = Font(name="Calibri", size=14, bold=True, color="003FB1")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="4B5563")
    section_font = Font(name="Calibri", size=12, bold=True, color="111827")
    bold_font = Font(name="Calibri", size=10, bold=True, color="111827")
    normal_font = Font(name="Calibri", size=10, color="374151")
    highlight_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    highlight_font = Font(name="Calibri", size=11, bold=True, color="92400E")

    ws_guide.cell(row=1, column=1, value="NEXUS CMS — Student Bulk Enrollment Specification").font = title_font
    ws_guide.cell(row=2, column=1, value="Suguna College of Engineering • Academic Governance & Enrollment Hub").font = subtitle_font

    ws_guide.cell(row=4, column=1, value="⭐ PREDEFINED PASSWORD FORMULA ENFORCEMENT").font = section_font
    ws_guide.merge_cells("A5:F5")
    rule_cell = ws_guide.cell(
        row=5,
        column=1,
        value="Default Password Rule: First 3 characters of student's name (lowercase only) + last 3 digits of register number.",
    )
    rule_cell.fill = highlight_fill
    rule_cell.font = highlight_font
    rule_cell.alignment = left_align
    ws_guide.row_dimensions[5].height = 26

    examples = [
        ("Student Name", "Register Number", "Name Prefix (3 chars lower)", "Reg Suffix (3 nums)", "Generated Default Password"),
        ("Arjun Patel", "714022AD001", "arj", "001", "arj001"),
        ("Abishek R K", "24AD012", "abi", "012", "abi012"),
        ("Bhavana S", "714022AD002", "bha", "002", "bha002"),
        ("Sneha K", "714022AD004", "sne", "004", "sne004"),
    ]

    for r_idx, row_data in enumerate(examples, start=7):
        ws_guide.row_dimensions[r_idx].height = 20
        for c_idx, val in enumerate(row_data, start=1):
            c = ws_guide.cell(row=r_idx, column=c_idx, value=val)
            c.border = thin_border
            if r_idx == 7:
                c.fill = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
                c.font = bold_font
                c.alignment = center_align
            else:
                c.font = normal_font
                c.alignment = center_align
                if c_idx == 5:
                    c.font = Font(name="Calibri", size=10, bold=True, color="003FB1")
                    c.fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")

    ws_guide.cell(row=14, column=1, value="COLUMN SPECIFICATIONS & RULES").font = section_font

    col_specs = [
        ("Column Name", "Required?", "Data Type", "Description & Rules"),
        ("register_number", "YES", "String", "Unique college registration/roll number (e.g. 714022AD001, 24AD012)"),
        ("full_name", "YES", "String", "Official student full name as per mark-sheet"),
        ("email", "Optional", "Email", "Institutional email. If omitted, generated as first_name.reg_no@suguna.edu.in"),
        ("phone", "Optional", "String", "Student active mobile phone number with +91 country code"),
        ("parent_name", "Optional", "String", "Parent or legal guardian full name"),
        ("parent_phone", "Optional", "String", "Parent/guardian emergency contact number"),
        ("gender", "Optional", "String", "Male / Female / Other"),
        ("date_of_birth", "Optional", "Date", "Format YYYY-MM-DD (e.g. 2005-04-18)"),
        ("blood_group", "Optional", "String", "O+, A+, B+, AB+, O-, A-, B-, AB-"),
        ("address", "Optional", "Text", "Permanent residential address"),
        ("department_code", "Optional", "String", "e.g. AIDS, CSE, ECE (Defaults to target department selected on screen)"),
        ("batch_name", "Optional", "String", "e.g. 2024-2028 (Defaults to target batch selected on screen)"),
        ("section_name", "Optional", "String", "e.g. A, B (Defaults to target section selected on screen)"),
    ]

    for r_idx, row_data in enumerate(col_specs, start=16):
        ws_guide.row_dimensions[r_idx].height = 20
        for c_idx, val in enumerate(row_data, start=1):
            c = ws_guide.cell(row=r_idx, column=c_idx, value=val)
            c.border = thin_border
            if r_idx == 16:
                c.fill = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
                c.font = bold_font
                c.alignment = center_align
            else:
                c.font = normal_font
                if c_idx == 2 and val == "YES":
                    c.font = Font(name="Calibri", size=10, bold=True, color="B91C1C")
                c.alignment = left_align

    ws_guide.column_dimensions["A"].width = 22
    ws_guide.column_dimensions["B"].width = 24
    ws_guide.column_dimensions["C"].width = 26
    ws_guide.column_dimensions["D"].width = 24
    ws_guide.column_dimensions["E"].width = 30
    ws_guide.column_dimensions["F"].width = 24

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
